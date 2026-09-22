from __future__ import annotations

import re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "plano-macro-zelogo.md"
OUTPUT = ROOT / "docs" / "Plano-Macro-Projeto-ZeloGO-v1.1.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text.strip())
    run.bold = bold
    run.font.size = Pt(8.2)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5A94")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_rich_text(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|\[[^\]]+\]\([^\)]+\)|`[^`]+`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^\)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(8.5)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    max_cols = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=max_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            set_repeat_table_header(row)
        for col_index in range(max_cols):
            value = values[col_index] if col_index < len(values) else ""
            set_cell_text(row.cells[col_index], value, bold=row_index == 0, color="FFFFFF" if row_index == 0 else None)
            if row_index == 0:
                set_cell_shading(row.cells[col_index], "1F4E79")
            elif row_index % 2 == 0:
                set_cell_shading(row.cells[col_index], "EEF3F7")
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip().strip("|")
        values = [cell.strip() for cell in raw.split("|")]
        if not all(re.fullmatch(r":?-+:?", value) for value in values):
            rows.append(values)
        index += 1
    return rows, index


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [("Title", 23, "17365D"), ("Heading 1", 16, "1F4E79"), ("Heading 2", 12.5, "2F75B5"), ("Heading 3", 10.5, "5B9BD5")]:
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Courier New"
        code_style.font.size = Pt(7.2)
        code_style.font.color.rgb = RGBColor.from_string("404040")
    header = section.header.paragraphs[0]
    header.text = "ZeloGO  |  Plano Macro do Projeto  |  Versão 1.1"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("6B7280")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("ZeloGO — documento executivo de construção, validação e escala  |  ")
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                paragraph = document.add_paragraph(style="Code Block")
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.space_after = Pt(6)
                paragraph.add_run("\n".join(code_lines))
                in_code = False
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if text == "Plano Macro do Projeto ZeloGO":
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run(text)
            else:
                paragraph = document.add_paragraph(style=f"Heading {level}")
                paragraph.add_run(text)
            index += 1
            continue
        image = re.fullmatch(r"!\[([^\]]*)\]\(([^\)]+)\)", stripped)
        if image:
            alt_text, relative_path = image.groups()
            image_path = SOURCE.parent / relative_path
            if image_path.exists():
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(image_path), width=Inches(6.7))
                paragraph.paragraph_format.space_after = Pt(4)
            else:
                paragraph = document.add_paragraph()
                paragraph.add_run(f"[Imagem não encontrada: {alt_text}]").italic = True
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        list_match = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if list_match:
            prefix, marker, text = list_match.groups()
            style = "List Number" if marker[0].isdigit() else "List Bullet"
            paragraph = document.add_paragraph(style=style)
            paragraph.paragraph_format.left_indent = Inches(0.25 + min(len(prefix), 4) * 0.08)
            add_rich_text(paragraph, text)
            index += 1
            continue
        if stripped.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.right_indent = Inches(0.3)
            run = paragraph.add_run(stripped[2:])
            run.italic = True
            run.font.color.rgb = RGBColor.from_string("4B5563")
            index += 1
            continue
        paragraph = document.add_paragraph()
        add_rich_text(paragraph, stripped)
        index += 1
    document.core_properties.title = "Plano Macro do Projeto ZeloGO — Versão 1.1"
    document.core_properties.subject = "Plano de construção, validação e escala de marketplace"
    document.core_properties.author = "Manus AI"
    document.core_properties.keywords = "ZeloGO, marketplace, PMBOK, Concierge MVP, unit economics, Trust and Safety"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
