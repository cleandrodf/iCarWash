from __future__ import annotations

from pathlib import Path
from textwrap import dedent

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "etapa-1" / "modelo-entrevistas-parceiros.xlsx"
MARKDOWN = ROOT / "docs" / "etapa-1" / "modelo-entrevistas-parceiros.md"

NAVY = "17324D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE = "F4F7FA"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
WHITE = "FFFFFF"
GRAY = "5B6573"
THIN_GRAY = Side(style="thin", color="D9E1F2")
BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)

PARTNERS = ["P1", "P2", "P3"]

QUESTIONS = [
    ("Q01", "Operação", "Quais serviços para carros vocês oferecem atualmente?"),
    ("Q02", "Operação", "Quais tipos de carro vocês atendem e quais serviços não realizam?"),
    ("Q03", "Operação", "Quantos atendimentos realizam aproximadamente em uma semana normal?"),
    ("Q04", "Aquisição", "Como chegam atualmente novos clientes?"),
    ("Q05", "Aquisição", "Qual foi o último novo cliente e como ele chegou até vocês?"),
    ("Q06", "Agenda", "Como a agenda é controlada e atualizada?"),
    ("Q07", "Agenda", "Quem atualiza a agenda e com que frequência?"),
    ("Q08", "Agenda", "Já ocorreu de um horário disponível não estar realmente livre? Por quê?"),
    ("Q09", "Demanda", "Quais dias e horários têm maior demanda?"),
    ("Q10", "Demanda", "Quais períodos normalmente ficam ociosos?"),
    ("Q11", "Disponibilidade", "Vocês conseguem indicar dois horários possíveis nos próximos 14 dias?"),
    ("Q12", "Oferta", "Quais são os três serviços mais procurados, seus preços e durações?"),
    ("Q13", "Oferta", "O preço muda conforme tamanho ou tipo do carro?"),
    ("Q14", "Oferta", "Há serviços adicionais contratados junto com a lavagem?"),
    ("Q15", "Oferta", "Existe valor mínimo para deslocamento ou atendimento fora do horário?"),
    ("Q16", "Qualificação", "O que precisa ser informado pelo cliente para confirmar o serviço?"),
    ("Q17", "Resposta", "Quanto tempo normalmente levam para responder a uma solicitação?"),
    ("Q18", "Resposta", "Por qual canal preferem receber uma nova solicitação?"),
    ("Q19", "Resposta", "Quem seria responsável por aceitar ou rejeitar o pedido?"),
    ("Q20", "Resposta", "Quais motivos fazem vocês rejeitarem uma solicitação?"),
    ("Q21", "Cancelamento", "Como tratam atrasos, cancelamentos e ausência do cliente?"),
    ("Q22", "Qualidade", "Como tratam retrabalho, reclamações ou serviço considerado incompleto?"),
    ("Q23", "Risco", "Já ocorreu dano ou reclamação relacionada ao veículo? Como foi tratada?"),
    ("Q24", "Piloto", "Aceitariam avaliar manualmente solicitações reais durante duas semanas?"),
    ("Q25", "Piloto", "Quais condições precisariam existir para o teste ser aceitável?"),
    ("Q26", "Piloto", "Quantos pedidos por semana conseguiriam avaliar sem prejudicar clientes atuais?"),
    ("Q27", "Piloto", "Quem seria o contato responsável durante o teste?"),
    ("Q28", "Piloto", "Qual seria o principal motivo para interromper a participação?"),
    ("Q29", "Valor", "O que precisaria acontecer para considerar o teste valioso?"),
    ("Q30", "Monetização", "Que tipo de demanda teria valor suficiente para justificar uma taxa?"),
    ("Q31", "Monetização", "Seria preferível pagar por solicitação, aceite ou serviço concluído? Por quê?"),
    ("Q32", "Segurança", "Quais controles são necessários para atendimento no endereço do cliente?"),
    ("Q33", "Dados", "Quais informações do cliente são realmente necessárias para aceitar o pedido?"),
]


def style_title(ws, title, subtitle, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, title)
    cell.font = Font(name="Aptos Display", size=16, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    sub = ws.cell(2, 1, subtitle)
    sub.font = Font(name="Aptos", size=10, italic=True, color=GRAY)
    sub.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 32


def style_header(row):
    for cell in row:
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def style_body(ws, start_row, end_row, start_col, end_col):
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=start_col, max_col=end_col):
        for cell in row:
            cell.font = Font(name="Aptos", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE)


def set_widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def add_table(ws, ref, name):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)


def add_validation(ws, formula, cell_range):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cell_range)


def build_workbook():
    wb = Workbook()
    wb.properties.title = "Modelo de entrevistas com parceiros — ZeloGO"
    wb.properties.subject = "Consolidação da descoberta de parceiros potenciais"
    wb.properties.creator = "ZeloGO"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    ws = wb.active
    ws.title = "Instrucoes"
    style_title(ws, "ZeloGO — modelo de entrevistas com parceiros", "Template para consolidar três entrevistas da Etapa 1. Não preencher dados pessoais desnecessários.", 4)
    instructions = [
        ("Objetivo", "Comparar três parceiros potenciais quanto a oferta, capacidade, agenda, riscos e disposição para o Concierge MVP."),
        ("Identificação", "Use P1, P2 e P3 como identificadores. Nomes e contatos devem permanecer em canal controlado, não no repositório."),
        ("Como preencher", "Registre fatos observados, exemplos recentes e fontes. Separe declaração futura de comportamento comprovado."),
        ("Nota de evidência", "0 = não observado; 1 = declaração genérica; 2 = exemplo parcial; 3 = evidência concreta e verificável."),
        ("Parceiro-design", "Só considerar quando houver capacidade verificável, disponibilidade concreta e disposição para avaliar solicitação real."),
        ("Privacidade", "Evite endereço completo, telefone, documentos e dados do veículo. Registre somente o mínimo necessário."),
        ("G0", "A meta da Etapa 1 é contatar três parceiros e obter pelo menos dois dispostos a avaliar o Concierge."),
        ("G1", "Os limiares do Gate G1 devem ser propostos antes do G0 e aprovados pelo comitê."),
    ]
    ws.cell(4, 1, "Tópico")
    ws.cell(4, 2, "Orientação")
    style_header(ws[4][0:2])
    for idx, (topic, text) in enumerate(instructions, 5):
        ws.cell(idx, 1, topic)
        ws.cell(idx, 2, text)
    style_body(ws, 5, 4 + len(instructions), 1, 2)
    set_widths(ws, {"A": 24, "B": 110, "C": 4, "D": 4})
    ws.freeze_panes = "A5"

    ws = wb.create_sheet("Cadastro")
    style_title(ws, "Cadastro das entrevistas", "Uma linha por parceiro potencial. Use identificadores P1, P2 e P3.", 12)
    headers = ["ID parceiro", "Tipo operação", "Região ampla", "Data entrevista", "Entrevistador", "Consentimento para notas", "Tempo operação", "Atendimentos/semana", "Canal autorizado", "Status entrevista", "Candidato design?", "Observação"]
    for col, value in enumerate(headers, 1):
        ws.cell(3, col, value)
    style_header(ws[3])
    for partner in PARTNERS:
        ws.append([partner, "", "", "", "", "", "", "", "", "Planejada", "", ""])
    style_body(ws, 4, 6, 1, len(headers))
    add_table(ws, "A3:L6", "CadastroParceiros")
    add_validation(ws, '"Estabelecimento,Prestador móvel,Híbrida"', "B4:B6")
    add_validation(ws, '"Sim,Não,Pendente"', "F4:F6")
    add_validation(ws, '"Planejada,Concluída,Reagendada,Cancelada"', "J4:J6")
    add_validation(ws, '"Sim,Não,Pendente"', "K4:K6")
    set_widths(ws, {"A": 12, "B": 20, "C": 20, "D": 16, "E": 18, "F": 24, "G": 18, "H": 20, "I": 24, "J": 18, "K": 18, "L": 36})
    ws.freeze_panes = "A4"

    ws = wb.create_sheet("Perguntas")
    style_title(ws, "Banco de perguntas", "Perguntas comuns aos três parceiros. Use Respostas para registrar a síntese de cada sessão.", 3)
    for col, value in enumerate(["ID pergunta", "Tema", "Pergunta"], 1):
        ws.cell(3, col, value)
    style_header(ws[3])
    for question in QUESTIONS:
        ws.append(list(question))
    style_body(ws, 4, 3 + len(QUESTIONS), 1, 3)
    add_table(ws, f"A3:C{3 + len(QUESTIONS)}", "BancoPerguntas")
    set_widths(ws, {"A": 14, "B": 20, "C": 100})
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:C{3 + len(QUESTIONS)}"

    ws = wb.create_sheet("Respostas")
    style_title(ws, "Respostas consolidadas", "Uma linha por parceiro e pergunta. Registre o resumo, o exemplo concreto e a implicação para o ZeloGO.", 12)
    response_headers = ["ID parceiro", "ID pergunta", "Tema", "Pergunta", "Resposta resumida", "Exemplo recente concreto", "Fonte/registro", "Nota evidência 0–3", "Dor ou objeção", "Hipótese afetada", "Implicação produto/operação", "Próxima ação"]
    for col, value in enumerate(response_headers, 1):
        ws.cell(3, col, value)
    style_header(ws[3])
    for partner in PARTNERS:
        for qid, topic, question in QUESTIONS:
            ws.append([partner, qid, topic, question, "", "", "", "", "", "", "", ""])
    last = 3 + len(PARTNERS) * len(QUESTIONS)
    style_body(ws, 4, last, 1, len(response_headers))
    add_table(ws, f"A3:L{last}", "RespostasParceiros")
    add_validation(ws, '"0,1,2,3"', f"H4:H{last}")
    ws.conditional_formatting.add(f"H4:H{last}", CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor=RED)))
    ws.conditional_formatting.add(f"H4:H{last}", CellIsRule(operator="equal", formula=["3"], fill=PatternFill("solid", fgColor=GREEN)))
    set_widths(ws, {"A": 12, "B": 14, "C": 18, "D": 60, "E": 48, "F": 48, "G": 24, "H": 16, "I": 34, "J": 30, "K": 42, "L": 32})
    ws.freeze_panes = "A4"

    ws = wb.create_sheet("Oferta")
    style_title(ws, "Oferta e capacidade", "Registre até três serviços por parceiro. Não trate cadastro sem disponibilidade como oferta viável.", 14)
    oferta_headers = ["ID parceiro", "Serviço", "Preço/regra", "Duração min", "Tipo carro", "Modalidade", "Região atendida", "Capacidade/dia", "Horários concretos", "Tempo resposta", "Restrição", "Fonte", "Nota capacidade 0–3", "Observação"]
    for col, value in enumerate(oferta_headers, 1):
        ws.cell(3, col, value)
    style_header(ws[3])
    for partner in PARTNERS:
        for _ in range(3):
            ws.append([partner, "", "", "", "", "", "", "", "", "", "", "", "", ""])
    last = 3 + len(PARTNERS) * 3
    style_body(ws, 4, last, 1, len(oferta_headers))
    add_table(ws, f"A3:N{last}", "OfertaParceiros")
    add_validation(ws, '"No estabelecimento,Móvel,Híbrida"', f"F4:F{last}")
    add_validation(ws, '"0,1,2,3"', f"M4:M{last}")
    set_widths(ws, {"A": 12, "B": 28, "C": 22, "D": 14, "E": 20, "F": 20, "G": 22, "H": 16, "I": 30, "J": 18, "K": 30, "L": 22, "M": 18, "N": 34})
    ws.freeze_panes = "A4"

    ws = wb.create_sheet("Piloto")
    style_title(ws, "Disposição para o Concierge MVP", "Consolide a resposta concreta do parceiro sobre um teste manual de duas semanas.", 15)
    piloto_headers = ["ID parceiro", "Aceita avaliar solicitações reais?", "Condição para participar", "Pedidos/semana suportados", "Contato/função", "Canal preferido", "Tempo máximo resposta", "Dois horários possíveis", "Critério de valor", "Motivo de interrupção", "Modelo de cobrança aceitável", "Pagamento no local entendido?", "Controle móvel necessário", "Próxima ação", "Estado"]
    for col, value in enumerate(piloto_headers, 1):
        ws.cell(3, col, value)
    style_header(ws[3])
    for partner in PARTNERS:
        ws.append([partner, "", "", "", "", "", "", "", "", "", "", "", "", "", "Pendente"])
    style_body(ws, 4, 6, 1, len(piloto_headers))
    add_table(ws, "A3:O6", "PilotoParceiros")
    add_validation(ws, '"Sim,Não,Condicional,Pendente"', "B4:B6")
    add_validation(ws, '"Sim,Não,Pendente"', "L4:L6")
    add_validation(ws, '"Pendente,Pronto para validação,Parceiro-design,Não priorizar"', "O4:O6")
    ws.conditional_formatting.add("B4:B6", FormulaRule(formula=['B4="Sim"'], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add("B4:B6", FormulaRule(formula=['B4="Não"'], fill=PatternFill("solid", fgColor=RED)))
    set_widths(ws, {"A": 12, "B": 24, "C": 38, "D": 22, "E": 22, "F": 22, "G": 22, "H": 30, "I": 36, "J": 32, "K": 30, "L": 24, "M": 34, "N": 30, "O": 24})
    ws.freeze_panes = "A4"

    ws = wb.create_sheet("Riscos")
    style_title(ws, "Riscos, segurança e dados", "Use esta aba para registrar riscos antes de qualquer atendimento real, especialmente no modo móvel.", 11)
    risco_headers = ["ID parceiro", "Risco ou situação", "Modalidade", "Dado necessário", "Compartilhamento de endereço", "Tratamento de dano/reclamação", "Questão jurídica/privacidade", "Severidade 0–3", "Controle proposto", "Responsável", "Estado"]
    for col, value in enumerate(risco_headers, 1):
        ws.cell(3, col, value)
    style_header(ws[3])
    for partner in PARTNERS:
        for _ in range(3):
            ws.append([partner, "", "", "", "", "", "", "", "", "", "Aberto"])
    last = 3 + len(PARTNERS) * 3
    style_body(ws, 4, last, 1, len(risco_headers))
    add_table(ws, f"A3:K{last}", "RiscosParceiros")
    add_validation(ws, '"No estabelecimento,Móvel,Híbrida"', f"C4:C{last}")
    add_validation(ws, '"0,1,2,3"', f"H4:H{last}")
    add_validation(ws, '"Aberto,Mitigado,Escalado,Encerrado"', f"K4:K{last}")
    set_widths(ws, {"A": 12, "B": 34, "C": 20, "D": 32, "E": 34, "F": 36, "G": 36, "H": 16, "I": 38, "J": 22, "K": 18})
    ws.freeze_panes = "A4"

    ws = wb.create_sheet("Comparativo")
    style_title(ws, "Síntese comparativa dos três parceiros", "Use a síntese para preparar a recomendação do G0; não substitui evidência registrada nas abas anteriores.", 18)
    comp_headers = ["ID parceiro", "Perguntas respondidas", "Exemplos concretos", "Nota média evidência", "Oferta com horário concreto", "Aceita piloto", "Disponibilidade 0–3", "Capacidade 0–3", "Resposta 0–3", "Adequação região 0–3", "Disposição piloto 0–3", "Segurança 0–3", "Valor percebido 0–3", "Total / 21", "Recomendação", "Principal valor", "Principal objeção", "Próxima decisão"]
    for col, value in enumerate(comp_headers, 1):
        ws.cell(3, col, value)
    style_header(ws[3])
    for row, partner in enumerate(PARTNERS, 4):
        ws.cell(row, 1, partner)
        ws.cell(row, 2, f'=COUNTIFS(Respostas!$A:$A,$A{row},Respostas!$E:$E,"<>")')
        ws.cell(row, 3, f'=COUNTIFS(Respostas!$A:$A,$A{row},Respostas!$F:$F,"<>")')
        ws.cell(row, 4, f'=IFERROR(AVERAGEIF(Respostas!$A:$A,$A{row},Respostas!$H:$H),"")')
        ws.cell(row, 5, f'=COUNTIFS(Oferta!$A:$A,$A{row},Oferta!$I:$I,"<>")')
        ws.cell(row, 6, f'=IFERROR(INDEX(Piloto!$B:$B,MATCH($A{row},Piloto!$A:$A,0)),"")')
        for col in range(7, 14):
            ws.cell(row, col, "")
        ws.cell(row, 14, f'=IF(COUNTA(G{row}:M{row})=0,"",SUM(G{row}:M{row}))')
        ws.cell(row, 15, f'=IF(N{row}="","",IF(N{row}>=15,"Priorizar para Concierge",IF(N{row}>=10,"Manter em validação","Não priorizar agora")))')
        ws.cell(row, 16, "")
        ws.cell(row, 17, "")
        ws.cell(row, 18, "")
    style_body(ws, 4, 6, 1, len(comp_headers))
    add_table(ws, "A3:R6", "ComparativoParceiros")
    add_validation(ws, '"0,1,2,3"', "G4:M6")
    ws.conditional_formatting.add("N4:N6", CellIsRule(operator="greaterThanOrEqual", formula=["15"], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add("N4:N6", CellIsRule(operator="between", formula=["10", "14"], fill=PatternFill("solid", fgColor=YELLOW)))
    ws.conditional_formatting.add("N4:N6", CellIsRule(operator="lessThan", formula=["10"], fill=PatternFill("solid", fgColor=RED)))
    set_widths(ws, {"A": 12, "B": 18, "C": 18, "D": 20, "E": 24, "F": 18, "G": 16, "H": 16, "I": 16, "J": 20, "K": 20, "L": 16, "M": 18, "N": 14, "O": 28, "P": 34, "Q": 34, "R": 34})
    ws.freeze_panes = "A4"

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.5
        sheet.page_margins.bottom = 0.5

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def build_markdown():
    lines = [
        "# Modelo de consolidação das entrevistas com parceiros — ZeloGO",
        "",
        "**Uso:** entrevistar três parceiros potenciais (P1, P2 e P3) durante a Etapa 1. **Escopo:** carros; motocicletas permanecem como hipótese posterior.",
        "",
        "> Registre fatos observados, exemplos recentes e disponibilidade concreta. Não trate uma declaração genérica de interesse como evidência de parceiro-design.",
        "",
        "## 1. Cadastro das sessões",
        "",
        "| Campo | P1 | P2 | P3 |",
        "|---|---|---|---|",
        "| Tipo de operação |  |  |  |",
        "| Região ampla |  |  |  |",
        "| Data da entrevista |  |  |  |",
        "| Consentimento para notas |  |  |  |",
        "| Status da entrevista | Planejada | Planejada | Planejada |",
        "| Candidato a parceiro-design? |  |  |  |",
        "",
        "## 2. Síntese das respostas",
        "",
        "| ID pergunta | Tema | P1 — resposta/evidência | P2 — resposta/evidência | P3 — resposta/evidência | Nota 0–3 | Implicação |",
        "|---|---|---|---|---|---:|---|",
    ]
    for qid, topic, question in QUESTIONS:
        lines.append(f"| {qid} | {topic} |  |  |  |  | {question} |")
    lines.extend([
        "",
        "## 3. Oferta e capacidade",
        "",
        "| Parceiro | Serviço | Preço/regra | Duração | Modalidade | Região | Horários concretos | Tempo de resposta | Restrição |",
        "|---|---|---|---:|---|---|---|---|---|",
        "| P1 |  |  |  |  |  |  |  |  |",
        "| P1 |  |  |  |  |  |  |  |  |",
        "| P1 |  |  |  |  |  |  |  |  |",
        "| P2 |  |  |  |  |  |  |  |  |",
        "| P2 |  |  |  |  |  |  |  |  |",
        "| P2 |  |  |  |  |  |  |  |  |",
        "| P3 |  |  |  |  |  |  |  |  |",
        "| P3 |  |  |  |  |  |  |  |  |",
        "| P3 |  |  |  |  |  |  |  |  |",
        "",
        "## 4. Disposição para o Concierge MVP",
        "",
        "| Campo | P1 | P2 | P3 |",
        "|---|---|---|---|",
        "| Aceita avaliar solicitações reais por duas semanas? |  |  |  |",
        "| Condições para participar |  |  |  |",
        "| Pedidos por semana suportados |  |  |  |",
        "| Contato/função |  |  |  |",
        "| Canal preferido |  |  |  |",
        "| Tempo máximo de resposta |  |  |  |",
        "| Dois horários possíveis nos próximos 14 dias |  |  |  |",
        "| Critério de valor |  |  |  |",
        "| Principal motivo para interromper |  |  |  |",
        "| Modelo de cobrança aceitável |  |  |  |",
        "| Próxima ação |  |  |  |",
        "",
        "## 5. Riscos e controles",
        "",
        "| Parceiro | Risco/situação | Modalidade | Dado necessário | Controle proposto | Severidade 0–3 | Responsável | Estado |",
        "|---|---|---|---|---|---:|---|---|",
        "| P1 |  |  |  |  |  |  | Aberto |",
        "| P2 |  |  |  |  |  |  | Aberto |",
        "| P3 |  |  |  |  |  |  | Aberto |",
        "",
        "## 6. Síntese para o G0",
        "",
        "| Critério | P1 | P2 | P3 | Observação do projeto |",
        "|---|---:|---:|---:|---|",
        "| Disponibilidade concreta 0–3 |  |  |  |  |",
        "| Capacidade verificável 0–3 |  |  |  |  |",
        "| Tempo de resposta 0–3 |  |  |  |  |",
        "| Adequação à microrregião 0–3 |  |  |  |  |",
        "| Disposição para piloto 0–3 |  |  |  |  |",
        "| Segurança e controles 0–3 |  |  |  |  |",
        "| Valor percebido 0–3 |  |  |  |  |",
        "| Total / 21 |  |  |  |  |",
        "| Recomendação |  |  |  | Priorizar, manter em validação ou não priorizar |",
        "",
        "## 7. Regra de decisão",
        "",
        "Para o G0, o projeto deve contatar três parceiros potenciais e buscar pelo menos dois dispostos a avaliar o Concierge MVP. Um parceiro só deve ser classificado como parceiro-design quando houver capacidade verificável, disponibilidade concreta e disposição para avaliar uma solicitação real.",
        "",
        "### Escala sugerida",
        "",
        "- **0:** não observado ou desconhecido;",
        "- **1:** declaração genérica, sem exemplo;",
        "- **2:** exemplo parcial ou evidência indireta;",
        "- **3:** evidência concreta, recente e verificável.",
        "",
        "## 8. Proteção de dados",
        "",
        "Use identificadores P1, P2 e P3. Não copie para o repositório nome completo, telefone, endereço residencial, documentos ou dados desnecessários do veículo. Mantenha informações de contato em canal controlado.",
    ])
    MARKDOWN.write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    build_workbook()
    build_markdown()
    print(f"Workbook: {OUTPUT}")
    print(f"Markdown: {MARKDOWN}")
